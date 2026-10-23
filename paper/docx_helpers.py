"""Small python-docx helpers for the paper: styles, numbered citations, booktabs-style tables, captions."""
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)


def _fonts(style_or_run_rpr, name=FONT):
    rfonts = style_or_run_rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        style_or_run_rpr.append(rfonts)
    for attr in list(rfonts.attrib):
        if "Theme" in attr:
            del rfonts.attrib[attr]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), name)


def set_style(style, size, bold=None, italic=None, name=FONT, before=0, after=6, line=1.12, align=None, keep_next=False):
    f = style.font
    f.name, f.size, f.color.rgb = name, Pt(size), BLACK
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    _fonts(style.element.get_or_add_rPr(), name)
    pf = style.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    pf.keep_with_next = keep_next
    if align is not None:
        pf.alignment = align


class Paper:
    def __init__(self, running_title):
        self.doc = Document()
        self.refs = {}      # key -> formatted reference
        self.order = []     # keys in order of first citation
        self.tables = self.figures = 0
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(1.0)
        st = self.doc.styles
        set_style(st["Normal"], 11, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        set_style(st["Heading 1"], 13.5, bold=True, before=14, after=6, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_style(st["Heading 2"], 11.5, bold=True, before=10, after=4, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_style(st["Heading 3"], 11, bold=True, italic=True, before=8, after=3, keep_next=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_style(st["Caption"], 9.5, bold=False, italic=False, before=4, after=8, line=1.05, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        # header and footer
        hp = s.header.paragraphs[0]
        hp.text = running_title
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in hp.runs:
            r.font.size, r.font.italic, r.font.name = Pt(9), True, FONT
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if t:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), t)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            run._r.append(el)
        run.font.size, run.font.name = Pt(9.5), FONT

    # ------------------------------------------------------------------ text with **bold**, ~~italic~~, [@cite]
    def cite_sub(self, s):
        def f(m):
            nums = []
            for k in m.group(1).split(","):
                k = k.strip().lstrip("@")
                if k not in self.refs:
                    raise KeyError(f"unknown reference key {k!r}")
                if k not in self.order:
                    self.order.append(k)
                nums.append(self.order.index(k) + 1)
            return "[" + ", ".join(str(n) for n in sorted(set(nums))) + "]"
        return re.sub(r"\[@([A-Za-z0-9_,@\s]+)\]", f, s)

    def runs(self, p, text, size=None, color=None):
        text = self.cite_sub(text)
        bold = ital = False
        for tok in re.split(r"(\*\*|~~)", text):
            if tok == "**":
                bold = not bold
            elif tok == "~~":
                ital = not ital
            elif tok:
                r = p.add_run(tok)
                r.bold, r.italic = bold or None, ital or None
                if size:
                    r.font.size = Pt(size)
                if color:
                    r.font.color.rgb = color
        return p

    def p(self, text, align=None, size=None, after=None, before=None, indent=None, keep_next=False):
        para = self.doc.add_paragraph()
        self.runs(para, text, size)
        if align is not None:
            para.alignment = align
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        if before is not None:
            para.paragraph_format.space_before = Pt(before)
        if indent:
            para.paragraph_format.left_indent = Inches(indent)
            para.paragraph_format.right_indent = Inches(indent)
        para.paragraph_format.keep_with_next = keep_next
        return para

    def h(self, text, level=1):
        return self.doc.add_heading(text, level=level)

    def bullets(self, items, size=None):
        for it in items:
            para = self.doc.add_paragraph(style="List Bullet")
            self.runs(para, it, size)
            para.paragraph_format.space_after = Pt(2)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def numbered(self, items):
        for it in items:
            para = self.doc.add_paragraph(style="List Number")
            self.runs(para, it)
            para.paragraph_format.space_after = Pt(2)

    def equation(self, text, number=None):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(3)
        para.paragraph_format.space_after = Pt(5)
        r = para.add_run(text)
        r.font.name, r.italic = "Cambria Math", True
        _fonts(r._r.get_or_add_rPr(), "Cambria Math")
        if number:
            n = para.add_run(f"        ({number})")
            n.font.name = FONT
        return para

    # ------------------------------------------------------------------ tables and figures
    @staticmethod
    def _border(cell, **edges):
        tcPr = cell._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders")
            tcPr.append(b)
        for edge, sz in edges.items():
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single" if sz else "nil")
            el.set(qn("w:sz"), str(sz or 0))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
            b.append(el)

    def table(self, caption, header, rows, widths, size=9, bold_last=False, notes=None, left=(0,), bold_rows=()):
        self.tables += 1
        cap = self.doc.add_paragraph(style="Caption")
        self.runs(cap, f"**Table {self.tables}.** {caption}")
        cap.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        data = [header] + rows
        for i, row in enumerate(data):
            tr = t.rows[i]
            trPr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            if i == 0:
                th = OxmlElement("w:tblHeader")
                trPr.append(th)
            for j, val in enumerate(row):
                c = tr.cells[j]
                c.width = Inches(widths[j])
                para = c.paragraphs[0]
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.space_before = Pt(1)
                para.paragraph_format.line_spacing = 1.0
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in left else WD_ALIGN_PARAGRAPH.CENTER
                is_b = i == 0 or (bold_last and i == len(data) - 1) or i - 1 in bold_rows
                self.runs(para, f"**{val}**" if is_b and "**" not in str(val) else str(val), size)
                edges = {}
                if i == 0:
                    edges.update(top=12, bottom=6)
                if i == len(data) - 1:
                    edges.update(bottom=12)
                if edges:
                    self._border(c, **edges)
        if notes:
            n = self.p(notes, size=8.5, after=8, before=2)
            n.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figure(self, path, caption, width=6.3):
        self.figures += 1
        self.doc.add_picture(path, width=Inches(width))
        pic = self.doc.paragraphs[-1]
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.keep_with_next = True
        pic.paragraph_format.space_after = Pt(2)
        cap = self.doc.add_paragraph(style="Caption")
        self.runs(cap, f"**Figure {self.figures}.** {caption}")

    def page_break(self):
        self.doc.add_page_break()

    def reference_list(self):
        for i, k in enumerate(self.order, 1):
            para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.35)
            para.paragraph_format.first_line_indent = Inches(-0.35)
            para.paragraph_format.space_after = Pt(3)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self.runs(para, f"[{i}]  " + self.refs[k], 9.5)

    def save(self, path):
        self.doc.save(path)
