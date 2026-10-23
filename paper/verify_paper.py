"""Quality checks on paper/mini_jev_paper.docx: structure, leftover markup, citation numbering, and text overlap.

    python paper/verify_paper.py

Overlap is measured against the repository's own README and report (self-overlap) as the share of the paper's
8-word sequences that also occur in them. It is not a substitute for a plagiarism service; it only shows how much of the
prose is reused from our own documents.
"""
import os
import re
import sys

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
doc = Document(os.path.join(HERE, "mini_jev_paper.docx"))

paras = [p.text for p in doc.paragraphs]
cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
text = "\n".join(paras + cells)
heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
print("headings:", len(heads), "| tables:", len(doc.tables), "| images:", len(doc.inline_shapes), "| paragraphs:", len(paras))
body_words = len(re.findall(r"\w+", "\n".join(paras)))
print("words in paragraphs:", body_words)

bad = [m for m in ("[@", "**", "~~", "TODO", "PLACEHOLDER", "{{", "XXX") if m in text]
print("leftover markup:", bad or "none")

refs_start = next(i for i, p in enumerate(paras) if p == "References")
ref_paras = [p for p in paras[refs_start + 1:] if re.match(r"^\[\d+\]", p)]
n_refs = len(ref_paras)
body = "\n".join(paras[:refs_start] + cells)
cited = set()
for m in re.finditer(r"\[(\d+(?:,\s*\d+)*)\]", body):
    cited |= {int(x) for x in m.group(1).split(",")}
print("references listed:", n_refs, "| cited numbers:", len(cited), "| max cited:", max(cited))
print("cited but not listed:", sorted(c for c in cited if c > n_refs) or "none")
print("listed but never cited in the body or tables:", sorted(set(range(1, n_refs + 1)) - cited) or "none")

# self-overlap with the repository's own documents
def grams(s, n=8):
    w = re.findall(r"[a-z0-9]+", s.lower())
    return {" ".join(w[i:i + n]) for i in range(len(w) - n + 1)}


own = ""
for f in ("README.md", os.path.join("docs", "REPORT.md"), os.path.join("docs", "METHODOLOGY.md")):
    own += open(os.path.join(ROOT, f), encoding="utf-8").read() + "\n"
paper_g = grams("\n".join(paras[:refs_start]))
own_g = grams(own)
print(f"8-gram overlap with README/REPORT/METHODOLOGY: {100 * len(paper_g & own_g) / max(1, len(paper_g)):.1f}% "
      f"({len(paper_g & own_g)} of {len(paper_g)})")
open(os.path.join(HERE, "_paper_text.txt"), "w", encoding="utf-8").write(text)
sys.exit(1 if bad else 0)
